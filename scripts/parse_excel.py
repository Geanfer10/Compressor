"""
Le a aba "Resumo" da planilha de compressores e gera os arquivos JSON
que o painel (index.html) consome.

Espera o arquivo em uploads/compressores.xlsx (sempre substituindo o
arquivo anterior com o mesmo nome, igual ja e feito no painel MTBF-MTTR).

Gera:
  data/latest.json        -> leitura mais recente
  data/AAAA-MM-DD.json    -> snapshot do dia
  data/history.json       -> historico (um registro por dia, mais recente por ultimo)
"""

import json
import os
from datetime import datetime, timezone

import openpyxl

EXCEL_PATH = "uploads/compressores.xlsx"
DATA_DIR = "data"
WARN_THRESHOLD = 1000  # horas restantes iguais ou abaixo disso entram em "atencao"


def norm(value):
    text = str(value or "").strip().lower()
    # remove acentos (NFD separa a letra do acento; filtramos as marcas combinantes)
    import unicodedata
    text = unicodedata.normalize("NFD", text)
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    return text


def find_header(rows):
    """Acha a linha e a coluna onde esta escrito 'Compressor'."""
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            if norm(cell) == "compressor":
                return i, j
    return None, None


def map_columns(header_row):
    cols = {}
    for j, cell in enumerate(header_row):
        n = norm(cell)
        if "ultima" in n and "prevent" in n:
            cols["last_date"] = j
        elif "carga" in n and "prevent" in n:
            cols["hrs_prev"] = j
        elif "atual" in n:
            cols["hrs_atual"] = j
        elif "diferen" in n:
            cols["diff"] = j
        elif "proxima" in n:
            cols["faltam"] = j
        elif "ordem" in n and "servi" in n:
            cols["ordem_servico"] = j
    return cols


def status_of(faltam):
    if faltam is None:
        return "sem_dados"
    if faltam < 0:
        return "atrasado"
    if faltam <= WARN_THRESHOLD:
        return "atencao"
    return "ok"


def as_number(value):
    return value if isinstance(value, (int, float)) else None


def main():
    if not os.path.exists(EXCEL_PATH):
        raise SystemExit(f"Arquivo nao encontrado: {EXCEL_PATH}")

    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    if "Resumo" not in wb.sheetnames:
        raise SystemExit("A aba 'Resumo' nao foi encontrada na planilha.")
    ws = wb["Resumo"]

    rows = list(ws.iter_rows(values_only=True))
    header_idx, compressor_col = find_header(rows)
    if header_idx is None:
        raise SystemExit("Nao encontrei a coluna 'Compressor' na aba Resumo.")

    cols = map_columns(rows[header_idx])

    compressores = []
    for row in rows[header_idx + 1:]:
        name = row[compressor_col]
        if name is None or str(name).strip() == "":
            break

        hrs_prev = as_number(row[cols["hrs_prev"]]) if "hrs_prev" in cols else None
        hrs_atual = as_number(row[cols["hrs_atual"]]) if "hrs_atual" in cols else None
        faltam = as_number(row[cols["faltam"]]) if "faltam" in cols else None
        diff = as_number(row[cols["diff"]]) if "diff" in cols else None
        last_date = row[cols["last_date"]] if "last_date" in cols else None

        if faltam is None and hrs_prev is not None and hrs_atual is not None:
            faltam = 4000 - (hrs_atual - hrs_prev)
        if diff is None and hrs_prev is not None and hrs_atual is not None:
            diff = hrs_atual - hrs_prev

        raw_os = row[cols["ordem_servico"]] if "ordem_servico" in cols else None
        ordem_servico = None
        if raw_os is not None and str(raw_os).strip() != "":
            ordem_servico = raw_os if isinstance(raw_os, (int, float)) else str(raw_os).strip()

        # Um compressor fora de uso e marcado escrevendo "Inativo" na coluna
        # "Ordem de Servico" da aba Resumo (em vez de um numero de OS). Isso
        # tira o compressor das contagens de status e do grafico do painel.
        status = "inativo" if isinstance(ordem_servico, str) and ordem_servico.lower() == "inativo" else status_of(faltam)

        compressores.append({
            "nome": str(name).strip(),
            "ultima_preventiva": last_date.isoformat() if hasattr(last_date, "isoformat") else None,
            "hrs_preventiva": hrs_prev,
            "hrs_atual": hrs_atual,
            "diferenca": diff,
            "faltam": faltam,
            "status": status,
            "ordem_servico": ordem_servico,
        })

    now = datetime.now(timezone.utc)
    payload = {
        "gerado_em": now.isoformat(),
        "compressores": compressores,
    }

    os.makedirs(DATA_DIR, exist_ok=True)

    with open(os.path.join(DATA_DIR, "latest.json"), "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)

    date_str = now.strftime("%Y-%m-%d")
    with open(os.path.join(DATA_DIR, f"{date_str}.json"), "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)

    history_path = os.path.join(DATA_DIR, "history.json")
    history = []
    if os.path.exists(history_path):
        with open(history_path, "r", encoding="utf-8") as f:
            history = json.load(f)

    history = [h for h in history if h.get("data") != date_str]
    history.append({"data": date_str, "compressores": compressores})
    history.sort(key=lambda h: h["data"])

    with open(history_path, "w", encoding="utf-8") as f:
        json.dump(history, f, ensure_ascii=False, indent=2)

    print(f"OK: {len(compressores)} compressores processados em {date_str}.")


if __name__ == "__main__":
    main()
